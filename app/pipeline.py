"""Runs the six-step workflow on one deck."""
from collections import defaultdict
from dataclasses import asdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from . import colors, concepts, config, deck, placement
from .icon_tools import insert_icon
from .library import Library
from .models import Flag, Report
from .selector import Selector

MODES = {"add", "suggest", "replace", "standardize"}
HERO_SHARE = 0.75   # an existing icon at least this share of the hero tier counts as a hero icon


def run(deck_path: Path, library: Library, mode: str = "add",
        slides: list[int] | None = None, out_path: Path | None = None) -> Report:
    if mode not in MODES:
        raise ValueError(f"mode must be one of {MODES}")
    if mode in {"replace", "standardize"}:
        raise NotImplementedError(f"{mode}: TODO (swap existing icons in place, keep xfrm)")

    prs = Presentation(deck_path)
    existing = deck.existing_icons(prs)                                  # 1. read
    settings = deck.detect_settings(prs, library, existing)
    items = deck.read_items(prs, slides)
    if mode == "add":
        on_slide = {e["slide"] for e in existing}
        items = [i for i in items if i.slide not in on_slide]          # TODO: per-item, not per-slide
    # Hero icons (section dividers, single-message slides) only when the deck already uses them: a theme-colored
    # icon at hero size on such a slide (style guide, Size and placement). Photos, badges and textures don't count.
    hero_flags = []
    hero_slides = {i.slide for i in deck.read_items(prs) if i.kind == "hero"}
    if not any(e["size"] >= HERO_SHARE * settings.tiers["hero"] and e.get("color_slot") and e["slide"] in hero_slides
               for e in existing):
        hero_flags = [Flag(i, "no hero icon: section and single-message slides get one only when the deck "
                              "already uses hero icons") for i in items if i.kind == "hero"]
        items = [i for i in items if i.kind != "hero"]

    # Each slide's candidates: its main component, then (if that can't take icons) its other sets of headings
    # or labels. Only these can get icons, so only these are sent to Claude.
    groups = defaultdict(list)
    for i in items:
        groups[i.set_id or f"solo-{id(i)}"].append(i)
    by_slide = defaultdict(list)
    for members in sorted(groups.values(), key=placement.priority):
        by_slide[members[0].slide].append(members)
    queue = {n: [m for k, m in enumerate(sets) if not config.ONE_ICON_SET_PER_SLIDE or k == 0
                 or m[0].rank <= config.FALLBACK_MAX_RANK] for n, sets in by_slide.items()}
    vocab = sorted(set(library.concept_map) | {t for i in library.icons.values() for t in i.tags})
    concept_notes = concepts.extract([i for sets in queue.values() for m in sets for i in m], vocab,
                                     list(library.concept_map))           # 2. concepts

    # 3-5a. Round by round, each slide still without icons offers its next candidate. Candidates with the same
    # structure (family) are decided together: icons on all or none, one placement. A set's icons count as
    # taken (same concept = same icon, one icon = one concept) only once placed, so a set that ends up without
    # icons doesn't keep an icon from a later set.
    sel = Selector(library, settings)
    sel.seed_ledger(existing)
    placer = placement.Placer(settings, prs)
    picks, flags, offered = [], list(hero_flags), set()
    while True:
        offer = [queue[n].pop(0) for n in sorted(queue) if queue[n]
                 and not (config.ONE_ICON_SET_PER_SLIDE and placer.taken(n))]
        if not offer:
            break
        offered |= {id(m) for m in offer}
        for family in _families(offer):
            family_picks, family_flags = _decide(family, sel, placer)
            picks += family_picks
            flags += family_flags
    for n, sets in by_slide.items():   # the rest: the slide's icons went to another set, or it got none
        for members in (m for m in sets if id(m) not in offered):
            taken = placer.taken(n)
            reason = placement.one_set_reason(taken) if taken else placement.main_only_reason(
                " ".join(sets[0][0].text.split())[:40])
            flags += [Flag(i, reason) for i in members]
    colors.assign(prs, picks, settings.color_slot)                       # 5b. theme color per icon

    report = Report(mode=mode, settings=settings, picks=picks, flags=flags, gaps=list(concept_notes))
    if library.indexed_from_filenames:
        report.gaps.append("library has no index: icons matched on file names only")

    if mode != "suggest":                                                # 5c. insert
        placements = [{"slide": p.item.slide, "icon_id": p.icon.id, "left": p.left, "top": p.top,
                       "size": p.size, "flip": p.flip, "color_slot": p.color_slot, "color_hex": p.color_hex,
                       "indent": p.indent, "shape_id": p.item.shape_id, "rtl": p.item.rtl} for p in picks]
        out, _ = insert(deck_path, library, placements, asdict(settings), out_path)
        report.output_file = str(out)
    return report                                                        # 6. report


def _families(offer: list[list]) -> list[list[list]]:
    """Group the offered sets by structure, at most one per slide; each family in the order it first appears."""
    if not config.SAME_STRUCTURE_ALL_OR_NONE:
        return [[m] for m in offer]
    families = {}
    for m in offer:
        key = m[0].family or f"solo-{id(m)}"
        while any(o[0].slide == m[0].slide for o in families.get(key, [])):
            key += "+"
        families.setdefault(key, []).append(m)
    return list(families.values())


def _decide(family: list[list], sel: Selector, placer) -> tuple[list, list[Flag]]:
    """Pick and place a family of sets (same structure, different slides): all of them get icons or none."""
    before = sel.snapshot()
    chosen, failed = [], {}
    for k, members in enumerate(family):
        set_picks, set_flags = sel.pick_set(members)
        if set_flags:
            failed[k] = set_flags
        chosen.append(set_picks)
    if not failed:
        placed = placer.place_family(chosen)
        failed = {k: f for k, f in enumerate(placed) if f}
    if not failed:
        return [p for c in chosen for p in c], []
    sel.restore(before)
    flags = []
    slides = sorted({m[0].slide for m in family})
    blocker = family[min(failed)][0].slide
    for k, members in enumerate(family):
        if k in failed:
            flags += failed[k]
        else:
            others = ", ".join(str(s) for s in slides if s != members[0].slide)
            flags += [Flag(i, f"held back to keep slides with the same structure consistent ({others}): slide "
                              f"{blocker} can't take icons on its matching set; see its reason",
                           [p.icon.id for p in chosen[k] if p.item is i]) for i in members]
    return [], flags


def insert(deck_path: Path, library: Library, placements: list[dict], settings: dict,
           out_path: Path | None = None) -> tuple[Path, list[str]]:
    """Insert icons at planned positions into a copy of the deck; the original is never overwritten.
    `placements`: dicts with slide, icon_id, left, top, size (inches), flip, and the icon's theme
    color_slot / color_hex (the deck color from `settings`, the report's settings dict, when absent).
    With indent (inches), the text of shape_id moves over by that much to make room (style guide, Rows).
    Returns the copy and the ids skipped because they are no longer approved."""
    prs = Presentation(deck_path)
    skipped = []
    for p in placements:
        icon = library.get(p["icon_id"])
        if icon is None:
            skipped.append(p["icon_id"])
            continue
        insert_icon(prs, p["slide"], library.svg_bytes(icon), p["left"], p["top"], p["size"],
                    p.get("color_hex") or settings["color_hex"], flip_h=p["flip"], name=f"Icon - {icon.id}",
                    theme_slot=p.get("color_slot") or settings.get("color_slot", "accent1"))
        if p.get("indent"):
            shape = colors.find_shape(prs.slides[p["slide"] - 1].shapes, p.get("shape_id"))
            if shape is not None and shape.has_text_frame:
                tf = shape.text_frame
                move = Inches(p["indent"] / getattr(shape, "scale", 1.0))   # in a scaled group's own units
                if p.get("rtl"):
                    tf.margin_right = tf.margin_right + move
                else:
                    tf.margin_left = tf.margin_left + move
    out_path = out_path or deck_path.with_name(deck_path.stem + "_icons.pptx")
    prs.save(out_path)
    return out_path, skipped


def placements_from_report(saved: dict, indices: list[int]) -> list[dict]:
    """Planned positions and colors for the chosen picks of a saved suggest report (report.to_dict)."""
    keys = ("left", "top", "size", "flip", "color_slot", "color_hex", "indent")
    return [{"slide": p["item"]["slide"], "icon_id": p["icon"]["id"], "shape_id": p["item"]["shape_id"],
             "rtl": p["item"].get("rtl", False), **{k: p.get(k) for k in keys}}
            for p in (saved["picks"][n] for n in indices)]
