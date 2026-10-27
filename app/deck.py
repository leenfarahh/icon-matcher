"""Step 1: read the deck. Extract items, existing icons and the deck's icon settings."""
import math
import re
from collections import Counter
from statistics import median

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import colors, config
from .icon_tools import EMU, THEME_CLASS
from .models import DeckSettings, Item

ARABIC = re.compile(r"[؀-ۿ]")
NUMBER_ONLY = re.compile(r"^\s*(\(?\d{1,3}[.)]?|[ivx]{1,4}[.)]|[a-z][.)])\s*$", re.I)   # page and step numbers
NOTE = re.compile(r"^\s*(sources?|notes?|disclaimer)\b", re.I)
SKIP_PLACEHOLDERS = ("TITLE", "FOOTER", "SLIDE_NUMBER", "DATE")   # TITLE also covers CENTER_TITLE and SUBTITLE
FOOTER_BAND = 0.92      # a box starting below this share of the slide height is a footer
SAME_SIZE = 0.85        # drawn shapes in one component: widths and heights within this ratio
SAME_WIDTH = 0.75       # plain text boxes in one component: widths within this ratio...
SAME_HEIGHT = 0.5       # ...and heights within this one (a header line is not its description box)
SAME_EDGE_IN = 0.15     # plain text boxes sharing a left or top edge this closely line up as one set
NODE_MAX_SHARE = 0.15   # a filled shape up to this share of the slide is a node or card, not a panel
LABEL_WORDS = 8         # a component of short labels (headings, node names) outranks longer text
DEFAULT_FONT_PT = 18.0  # PowerPoint's default text size
AGREE = 2 / 3           # share of existing icons that must use one color for the deck to adopt it
MIN_ICON_IN = 0.15      # anything smaller is a dot or a bullet, not an icon
# Icon holders: empty circles, rounded squares or badges the designer left for icons, two or more alike.
HOLDER_GEOMS = {"ellipse", "flowChartConnector", "donut", "roundRect", "hexagon", "octagon", "diamond"}
HOLDER_PARTS = HOLDER_GEOMS | {"rect", "arc", "blockArc", "pie", "chord", "snip1Rect", "round1Rect", "round2SameRect"}
HOLDER_MIN_IN, HOLDER_MAX_IN = 0.3, 1.6
HOLDER_ICON_SHARE = 0.6     # an icon fills this share of its holder's inner circle or square
HOLDER_PANEL_SHARE = 0.25   # a shape the holder sits in, up to this share of the slide, is the holder's panel
HOLDER_REACH_IN = 1.0       # with no panel, a label this close in line with the holder is its label
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _font_pt(shape, para=None) -> float:
    for p in [para] if para is not None else shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return DEFAULT_FONT_PT if para is None else _font_pt(shape)  # TODO: inherited placeholder sizes


def _bold(shape) -> bool:
    return next((bool(r.font.bold) for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip()), False)


def _not_an_item(shape, slide_h) -> bool:
    """Titles, footers, page and step numbers, and source notes never get icons."""
    if shape.is_placeholder and any(k in str(shape.placeholder_format.type) for k in SKIP_PLACEHOLDERS):
        return True
    text = shape.text_frame.text.strip()
    return bool(NUMBER_ONLY.match(text) or NOTE.match(text)) or shape.top / EMU > slide_h * FOOTER_BAND


def _marked(shape, p) -> bool:
    """The paragraph carries a bullet or numbering (its own, or inherited by a body placeholder)."""
    ppr = p._p.find(A + "pPr")
    if ppr is not None and ppr.find(A + "buNone") is not None:
        return False
    if ppr is not None and (ppr.find(A + "buChar") is not None or ppr.find(A + "buAutoNum") is not None):
        return True
    return shape.is_placeholder and any(k in str(shape.placeholder_format.type) for k in ("BODY", "OBJECT"))


def _bulleted(shape, paras) -> bool:
    """A real list: 2+ paragraphs with bullets or numbering. A heading line over a description is one card."""
    return sum(_marked(shape, p) for p in paras) >= 2


NARROW, SEMI, WIDE = set("ijl.,:;'!|`"), set("ftrI()[]{}-/\\\""), set("mwMW@%")


def _char_em(c: str) -> float:
    """Approximate width of a character in ems (proportional sans-serif such as Arial or Helvetica)."""
    if c == " ":
        return 0.28
    if c in NARROW:
        return 0.25
    if c in SEMI:
        return 0.34
    if c in WIDE:
        return 0.85
    if c.isdigit():
        return 0.56
    if c.isupper():
        return 0.68
    if c.islower():
        return 0.52
    return 0.6


def text_width(text: str, font_pt: float, bold: bool = False) -> float:
    """Estimated width of a run of text in inches (bold is about 7% wider)."""
    return sum(_char_em(c) for c in text) * font_pt / 72 * (1.07 if bold else 1.0)


def _line_widths(text: str, width_in: float, font_pt: float, bold: bool) -> list[float]:
    """Width in inches of each line when text wraps at word boundaries in width_in."""
    space = text_width(" ", font_pt, bold)
    out, used = [], 0.0
    for word in text.split():
        w = text_width(word, font_pt, bold)
        if used and used + space + w > width_in:
            out.append(used)
            used = 0.0
        if w > width_in:                       # a word longer than a line breaks across lines
            out += [width_in] * int(w // max(width_in, 1e-6))
            used = w % max(width_in, 1e-6) or width_in
        else:
            used += (space if used else 0.0) + w
    return out + [used]


def _wrap(text: str, width_in: float, font_pt: float, bold: bool) -> tuple[int, float]:
    """(lines, width of the first line in inches) when text wraps at word boundaries in width_in."""
    widths = _line_widths(text, width_in, font_pt, bold)
    return len(widths), widths[0]


def _widths(text: str, width_in: float, font_pt: float, bold: bool = False, wraps: bool = True) -> list[float]:
    """Width of each line of a paragraph, honoring manual line breaks (\\v)."""
    parts = [s.strip() for s in text.split("\v")] or [""]
    if not wraps:
        return [text_width(s, font_pt, bold) for s in parts]
    return [w for s in parts for w in _line_widths(s, width_in, font_pt, bold)]


def _lines(text: str, width_in: float, font_pt: float, bold: bool = False, wraps: bool = True) -> tuple[int, float]:
    """(lines, width of the first line) of a paragraph, honoring manual line breaks (\\v)."""
    widths = _widths(text, width_in, font_pt, bold, wraps)
    return len(widths), widths[0]


def text_lines(text: str, width_in: float, font_pt: float, bold: bool = False) -> int:
    """Estimated wrapped lines of a paragraph in a box `width_in` wide (insets included)."""
    return _lines(text, max(0.3, width_in - 2 * config.TEXT_INSET_IN), font_pt, bold)[0]


def visual_box(shape) -> tuple[float, float, float, float]:
    """(left, top, width, height) in inches of the area the shape covers on screen (rotation applied)."""
    return colors.rotated_box(shape.left / EMU, shape.top / EMU, shape.width / EMU, shape.height / EMU,
                              getattr(shape, "rotation", 0) or 0)


def text_layout(shape, drawn: bool) -> list[dict]:
    """Where each non-empty paragraph's first line sits, in inches: left and right ends of the line, its top,
    line height, alignment, and how far the text could move over without wrapping onto more lines; and "rows",
    every line of the shape's text as (left, top, right, bottom).
    Estimated from insets, vertical anchor, alignment (left when unset, as in PowerPoint), indent, line breaks,
    character widths and word wrapping. A rotated shape, or one whose text is turned within it, is taken to
    show its label upright and centered in the area it covers (its own top, left and insets point elsewhere)."""
    l, t, w, h = visual_box(shape)
    body = shape._element.find(".//" + A + "bodyPr")
    turned = bool((getattr(shape, "rotation", 0) or 0) % 180) or \
        (body is not None and body.get("vert") not in (None, "horz"))
    def inset(key, default):
        return 0.05 if turned else (int(body.get(key)) / EMU if body is not None and body.get(key) else default)
    li, ti, ri, bi = inset("lIns", 0.1), inset("tIns", 0.05), inset("rIns", 0.1), inset("bIns", 0.05)
    if not turned and colors.geometry(shape) == "chevron":   # a chevron's text starts after its notch
        li += 0.5 * min(w, h)
    cnv = shape._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvSpPr")
    is_text_box = cnv is not None and cnv.get("txBox") == "1"
    anchor = "ctr" if turned else \
        (body.get("anchor") if body is not None else None) or ("t" if is_text_box or not drawn else "ctr")
    default_align = shape._element.find(f".//{A}lstStyle/{A}lvl1pPr")
    default_align = (default_align.get("algn") if default_align is not None else None) or "l"
    wraps = body is None or body.get("wrap") != "none"
    bold = _bold(shape)

    paras = []
    for p in (p for p in shape.text_frame.paragraphs if p.text.strip()):
        ppr = p._p.find(A + "pPr")
        mar = int(ppr.get("marL", "0")) / EMU if ppr is not None else 0.0
        font = _font_pt(shape, p)
        widths = _widths(p.text.strip(), max(0.3, w - li - ri - mar), font, bold, wraps)
        align = "ctr" if turned else (ppr.get("algn") if ppr is not None else None) or default_align
        paras.append({"text": p.text.strip(), "font": font, "mar": mar, "lines": len(widths), "first": widths[0],
                      "widths": widths, "line_h": font * 1.2 / 72, "align": align})
    total = sum(p["lines"] for p in paras)
    room = 0.0   # how far the text could move over (inches) and still take no more lines
    for step in (x / 20 for x in range(1, 25)):
        if not wraps or sum(_lines(p["text"], max(0.3, w - li - ri - p["mar"] - step), p["font"], bold)[0]
                            for p in paras) <= total:
            room = step
        else:
            break
    block = sum(p["lines"] * p["line_h"] for p in paras)
    top = t + ti if anchor == "t" else (t + h - bi - block if anchor == "b" else t + (h - block) / 2)
    out, rows = [], []
    for p in paras:
        inner_l, inner_r = l + li + p["mar"], l + w - ri
        for k, line in enumerate(p["widths"]):
            width = min(inner_r - inner_l, line)
            x0 = {"ctr": (inner_l + inner_r - width) / 2, "r": inner_r - width}.get(p["align"], inner_l)
            y = top + k * p["line_h"]
            rows.append((x0, y, x0 + width, y + p["line_h"]))
            if k == 0:   # the first line, as wrapped
                out.append({"tx_left": x0, "tx_right": x0 + width, "tx_top": top, "line_h": p["line_h"],
                            "align": "ctr" if p["align"] == "ctr" else ("r" if p["align"] == "r" else "l"),
                            "room": room, "rows": rows})
        top += p["lines"] * p["line_h"]
    return out


def _nodes(theme, slide) -> list:
    """(area, (left, top, width, height), shape, fill hex) of every filled shape small enough to be a node,
    tile or card (not a panel or background), in inches."""
    slide_area = (slide.part.package.presentation_part.presentation.slide_width / EMU) * \
                 (slide.part.package.presentation_part.presentation.slide_height / EMU)
    out = []
    for (l, t, r, b), fill, sh in colors.boxes(theme, slide.shapes):
        l, t, r, b = l / EMU, t / EMU, r / EMU, b / EMU
        if isinstance(fill, tuple) and (r - l) * (b - t) <= NODE_MAX_SHARE * slide_area:
            out.append(((r - l) * (b - t), (l, t, r - l, b - t), sh, fill[0]))
    return out


def _style(shape, theme, nodes) -> dict:
    """What makes boxes one component. Text drawn in its own shape, or sitting on a node-sized filled shape,
    takes that shape's outline, fill and size (a hexagon labelled directly and one with a text box laid on it
    are the same node); plain text takes its text style and how it lines up with its neighbors."""
    fill = colors.shape_fill(theme, shape)
    own = colors.is_drawn(shape, fill)
    vb = visual_box(shape)
    center = (vb[0] + vb[2] / 2, vb[1] + vb[3] / 2)
    node = None if own else min(((a, box, sh, hexv) for a, box, sh, hexv in nodes
                                 if sh._element is not shape._element and box[0] <= center[0] <= box[0] + box[2]
                                 and box[1] <= center[1] <= box[1] + box[3]), default=None, key=lambda n: n[0])
    if own:
        box, geom, hexv = vb, colors.geometry(shape), fill[0] if isinstance(fill, tuple) else fill
    elif node:
        box, geom, hexv = node[1], colors.geometry(node[2]), node[3]
    else:
        box, geom, hexv = vb, colors.geometry(shape), None
    paras = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    return {"drawn": own or node is not None, "overlay": node is not None, "geom": geom, "fill": hexv,
            "font": _font_pt(shape), "bold": _bold(shape), "paras": len(paras), "box": box,
            "tw": vb[2], "th": vb[3], "left": vb[0], "top": vb[1],
            "own_drawn": own}


def _ratio(a, b) -> float:
    return min(a, b) / max(a, b) if max(a, b) > 0 else 1.0


def _alike(a, b) -> bool:
    if (a["drawn"], a["bold"]) != (b["drawn"], b["bold"]) or abs(a["font"] - b["font"]) > 0.5:
        return False
    if a["drawn"]:   # nodes, tiles, cards: same outline, fill and size wherever they sit (a ring, a grid, a row)
        ok = a["geom"] == b["geom"] and a["fill"] == b["fill"] and _ratio(a["box"][2], b["box"][2]) >= SAME_SIZE \
            and _ratio(a["box"][3], b["box"][3]) >= SAME_SIZE
        return ok and (not (a["overlay"] and b["overlay"]) or _ratio(a["th"], b["th"]) >= SAME_HEIGHT)
    lined_up = abs(a["left"] - b["left"]) < SAME_EDGE_IN or abs(a["top"] - b["top"]) < SAME_EDGE_IN
    return a["geom"] == b["geom"] and a["paras"] == b["paras"] and _ratio(a["th"], b["th"]) >= SAME_HEIGHT \
        and (_ratio(a["tw"], b["tw"]) >= SAME_WIDTH or lined_up)


def _alike_lists(a, b) -> bool:
    """Bullet lists of one component (the bullets under each column): same text style and outline, similar
    widths or lined up. Their lengths may differ."""
    if (a["drawn"], a["bold"], a["geom"]) != (b["drawn"], b["bold"], b["geom"]) or abs(a["font"] - b["font"]) > 0.5:
        return False
    lined_up = abs(a["left"] - b["left"]) < SAME_EDGE_IN or abs(a["top"] - b["top"]) < SAME_EDGE_IN
    return _ratio(a["tw"], b["tw"]) >= SAME_WIDTH or lined_up


def _components(boxes, theme, nodes, lists=False, keep_singles=False) -> list[list]:
    """Groups of 2+ boxes that look alike, in reading order: one component (all hexagons of a diagram, all
    headings of a set of columns, the bullet lists under each column), however they are arranged."""
    alike = _alike_lists if lists else _alike
    groups = []
    for b in sorted(boxes, key=lambda s: (s.top, s.left)):
        st = _style(b, theme, nodes)
        g = next((g for g in groups if alike(g[0][1], st)), None)
        if g:
            g.append((b, st))
        else:
            groups.append([(b, st)])
    return [g for g in groups if keep_singles or len(g) >= 2]


def _slide_title(slide) -> str:
    for s in slide.placeholders:
        if "TITLE" in str(s.placeholder_format.type) and "SUB" not in str(s.placeholder_format.type) \
                and s.has_text_frame and s.text_frame.text.strip():
            return " ".join(s.text_frame.text.split())
    return ""


def _inches(sh) -> tuple[float, float, float, float]:
    l, t, w, h = visual_box(sh)
    return l, t, l + w, t + h


def _holder_parts(sh) -> list | None:
    """The text-free preset shapes (circles, rings, rounded squares) a holder is drawn with: the shape itself,
    or all of a group's shapes. None when it has text, a picture, a freeform (an icon) or anything else."""
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        parts = []
        for c in sh.shapes:
            sub = _holder_parts(c)
            if sub is None:
                return None
            parts += sub
        return parts or None
    if not sh._element.tag.endswith("}sp") or (sh.has_text_frame and sh.text_frame.text.strip()):
        return None
    return [sh] if colors.geometry(sh) in HOLDER_PARTS else None


def _ids(sh) -> set[int]:
    return {int(e.get("id")) for e in sh._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")}


def holders(slide, theme, slide_h) -> list[list[dict]]:
    """Empty icon holders, in sets of alike ones (two or more), each paired with its label: the text in the
    panel the holder sits in, else the nearest text in line with it. A set is dropped when a holder has no label
    or two share one. Each holder: {"box", "ids", "label" (shape), "seat" (center x, center y, icon size)}."""
    found, inside = [], set()
    for sh in colors.leaves(slide.shapes, groups=True):
        if sh.shape_id in inside:
            continue
        l, t, r, b = _inches(sh)
        w, h = r - l, b - t
        parts = _holder_parts(sh) if min(w, h) >= HOLDER_MIN_IN and max(w, h) <= HOLDER_MAX_IN \
            and _ratio(w, h) >= 0.8 else None
        if not parts or (sh.shape_type != MSO_SHAPE_TYPE.GROUP and colors.geometry(sh) not in HOLDER_GEOMS):
            continue
        if not any(colors.is_drawn(p, colors.shape_fill(theme, p)) for p in parts):
            continue
        ids = _ids(sh)
        inside |= ids
        found.append({"box": (l, t, r, b), "ids": ids, "shape": sh,
                      "kind": "/".join(sorted(colors.geometry(p) for p in parts))})
    if not found:
        return []
    placed = {s.shape_id: s for s in colors.leaves(slide.shapes)}   # where each shape sits on the slide
    others = [s for s in placed.values() if not any(s.shape_id in f["ids"] for f in found)]
    texts = [s for s in others if s.has_text_frame and s.text_frame.text.strip() and not _not_an_item(s, slide_h)]
    slide_area = slide.part.package.presentation_part.presentation.slide_width / EMU * slide_h

    def overlaps(a, b):
        return min(a[2], b[2]) - max(a[0], b[0]) > 0.03 and min(a[3], b[3]) - max(a[1], b[1]) > 0.03

    rows = {s.shape_id: text_layout(s, colors.is_drawn(s, colors.shape_fill(theme, s)))[0]["rows"] for s in texts}

    def empty(f):   # nothing on it: no text over it (a panel's text elsewhere is fine), no picture or shape within it
        l, t, r, b = f["box"]
        for s in others:
            sl, st, sr, sb = _inches(s)
            within = (sr - sl) * (sb - st) < (r - l) * (b - t) and l <= (sl + sr) / 2 <= r and t <= (st + sb) / 2 <= b
            if within or any(overlaps(row, f["box"]) for row in rows.get(s.shape_id, ())):
                return False
        return True

    def seat(f):   # the icon goes in the innermost filled circle or square, else the holder itself
        l, t, r, b = f["box"]
        cx, cy = (l + r) / 2, (t + b) / 2
        inner = [_inches(placed[p.shape_id]) for p in _holder_parts(f["shape"]) if colors.geometry(p) in HOLDER_GEOMS
                 and colors.shape_fill(theme, p) is not None and p.shape_id in placed]
        inner = [(il, it, ir, ib) for il, it, ir, ib in inner if il <= cx <= ir and it <= cy <= ib]
        side = min([min(ir - il, ib - it) for il, it, ir, ib in inner] + [min(r - l, b - t)])
        return round(cx, 3), round(cy, 3), round(HOLDER_ICON_SHARE * side, 3)

    def center(s):
        sl, st, sr, sb = _inches(s)
        return (sl + sr) / 2, (st + sb) / 2

    def label(f):
        l, t, r, b = f["box"]
        cx, cy = (l + r) / 2, (t + b) / 2
        panels = [(s, _inches(s)) for s in others if colors.is_drawn(s, colors.shape_fill(theme, s))]
        panels = [(pl, pt, pr, pb) for s, (pl, pt, pr, pb) in panels if pl <= cx <= pr and pt <= cy <= pb
                  and 2 * (r - l) * (b - t) <= (pr - pl) * (pb - pt) <= HOLDER_PANEL_SHARE * slide_area]
        if panels:
            pl, pt, pr, pb = min(panels, key=lambda e: (e[2] - e[0]) * (e[3] - e[1]))
            on = [s for s in texts if pl <= center(s)[0] <= pr and pt <= center(s)[1] <= pb]
            if on:   # the panel's heading: its topmost text
                return min(on, key=lambda s: (_inches(s)[1], len(s.text_frame.text.split())))
        near = []
        for s in texts:
            sl, st, sr, sb = _inches(s)
            if min(r, sr) - max(l, sl) > 0:          # in its column: above or below it
                gap = max(st - b, t - sb)
            elif min(b, sb) - max(t, st) > 0:        # in its row: beside it
                gap = max(sl - r, l - sr)
            else:
                continue
            if -0.05 <= gap <= HOLDER_REACH_IN:
                near.append((gap, s))
        return min(near, key=lambda e: e[0])[1] if near else None

    groups = []
    for f in (f for f in found if empty(f)):
        g = next((g for g in groups if g[0]["kind"] == f["kind"] and _ratio(g[0]["box"][2] - g[0]["box"][0],
                                                                         f["box"][2] - f["box"][0]) >= SAME_SIZE), None)
        if g:
            g.append(f)
        else:
            groups.append([f])
    out = []
    for g in (g for g in groups if len(g) >= 2):
        for f in g:
            f["label"], f["seat"] = label(f), seat(f)
        labels = [f["label"].shape_id for f in g if f["label"] is not None]
        if len(labels) == len(g) == len(set(labels)):
            out.append(sorted(g, key=lambda f: (round(_inches(f["label"])[1], 1), _inches(f["label"])[0])))
    return out


def _family(slide, kind, st, member_boxes, slide_w, slide_h, extra="") -> str:
    """A set's structure, to spot the same component on other slides: the slide's layout, the set's kind and
    style, and the part of the slide (a 3x3 grid) its center falls in."""
    l, t = min(b[0] for b in member_boxes), min(b[1] for b in member_boxes)
    r, bottom = max(b[0] + b[2] for b in member_boxes), max(b[1] + b[3] for b in member_boxes)
    col = min(2, max(0, int((l + r) / 2 / slide_w * 3)))
    row = min(2, max(0, int((t + bottom) / 2 / slide_h * 3)))
    return "|".join([slide.slide_layout.name, kind, str(st["drawn"]), st["geom"], str(st["fill"]),
                     f"{st['font']:.0f}", str(st["bold"]), f"{col}{row}", extra])


def read_items(prs: Presentation, slides: list[int] | None = None) -> list[Item]:
    """Icon candidates, in sets that get icons together or not at all. Text inside groups counts too.
    rank: -1 the labels of empty icon holders; 0 nodes, tiles and cards with short labels, and heroes; 1 short
    plain labels; 2 longer text; 3 bullet lists. The best-ranked set is the slide's main component."""
    items: list[Item] = []
    slide_w, slide_h = prs.slide_width / EMU, prs.slide_height / EMU
    themes = {}
    for n, slide in enumerate(prs.slides, 1):
        if slides and n not in slides:
            continue
        master = slide.slide_layout.slide_master
        theme = themes.setdefault(id(master), colors.Theme(master))
        title = _slide_title(slide)
        nodes = _nodes(theme, slide)
        boxes = [s for s in colors.leaves(slide.shapes) if s.has_text_frame and s.text_frame.text.strip()
                 and not _not_an_item(s, slide_h)]
        # Empty holders the designer left for icons: the icon goes in the holder and stands for its label.
        for k, group in enumerate(holders(slide, theme, slide_h), 1):
            style = {"drawn": True, "geom": group[0]["kind"], "fill": None, "font": 0, "bold": False}
            fam = _family(slide, "holder", style, [visual_box(f["shape"]) for f in group], slide_w, slide_h,
                          f"{group[0]['box'][2] - group[0]['box'][0]:.1f}")
            for f in group:
                st = _style(f["label"], theme, nodes)
                it = _item(n, f["label"], "card", f"s{n}-holders{k}", text_layout(f["label"], st["own_drawn"])[0],
                           -1, title, st)
                it.holder, it.family = f["seat"], fam
                items.append(it)
        others, lists = [], []
        for b in boxes:   # one idea per box (a heading, or heading + text), or a bulleted list
            paras = [p for p in b.text_frame.paragraphs if p.text.strip()]
            (lists if len(paras) >= 2 and _bulleted(b, paras) else others).append(b)
        # Bullet lists: an item per bulleted paragraph (not a heading line above them), on its first line.
        # Lists that look alike (the bullets under each column) are one set: all of them get icons, or none.
        for k, group in enumerate(_components(lists, theme, nodes, lists=True, keep_singles=True), 1):
            fam = _family(slide, "list", group[0][1], [visual_box(b) for b, _ in group], slide_w, slide_h)
            for b, st in group:
                line = 0
                paras = [p for p in b.text_frame.paragraphs if p.text.strip()]
                for i, (p, geo) in enumerate(zip(paras, text_layout(b, st["own_drawn"]))):
                    if _marked(b, p):
                        it = _item(n, b, "list", f"s{n}-list{k}", geo, 3, title, st, text=p.text.strip())
                        it.paragraph, it.line, it.family = i, line, fam
                        items.append(it)
                    line += text_lines(p.text.strip(), visual_box(b)[2], _font_pt(b, p))
        for k, group in enumerate(_components(others, theme, nodes), 1):
            words = median(len(b.text_frame.text.split()) for b, _ in group)
            short = words <= LABEL_WORDS
            rank = (0 if short else 2) if group[0][1]["drawn"] else (1 if short else 2)
            fam = _family(slide, "card", group[0][1], [st["box"] for _, st in group], slide_w, slide_h)
            for b, st in group:
                it = _item(n, b, "card", f"s{n}-set{k}", text_layout(b, st["own_drawn"])[0], rank, title, st)
                it.family = fam
                items.append(it)
        # A lone short message on a slide is a hero (added only when the deck already uses hero icons).
        if len(boxes) == 1 and others and len(others[0].text_frame.text) < 120:
            st = _style(others[0], theme, nodes)
            items.append(_item(n, others[0], "hero", None, text_layout(others[0], st["own_drawn"])[0], 0, title, st))
        # TODO: KPI tiles, SmartArt, tables
    return items


def _item(n, shape, kind, set_id, geo, rank, title, st, text=None) -> Item:
    """An item's box is the node or card its text is drawn in or sits on, else the text box itself."""
    t = text or shape.text_frame.text.strip()
    left, top, w, h = st["box"]
    return Item(slide=n, shape_id=shape.shape_id, text=t, kind=kind, left=left, top=top, w=w, h=h,
                font_pt=_font_pt(shape), set_id=set_id, rtl=bool(ARABIC.search(t)), rank=rank, slide_title=title,
                in_shape=st["drawn"], room=geo["room"], rows=[tuple(round(v, 3) for v in r) for r in geo["rows"]],
                **{k: geo[k] for k in ("tx_left", "tx_right", "tx_top", "line_h", "align")})


def existing_icons(prs: Presentation) -> list[dict]:
    """Icon-like shapes, with the theme slot they're drawn in when it can be read (else color_slot None).
    Named 'Icon - <id>' when this app (or the skill) placed them. Empty icon holders and dots don't count."""
    found, themes = [], {}
    slide_h = prs.slide_height / EMU
    for n, slide in enumerate(prs.slides, 1):
        master = slide.slide_layout.slide_master
        theme = themes.setdefault(id(master), colors.Theme(master))
        empty = {i for group in holders(slide, theme, slide_h) for f in group for i in f["ids"]}
        for s in slide.shapes:
            if s.left is None or (s.has_text_frame and s.text_frame.text.strip()) or s.shape_id in empty:
                continue
            w, h = s.width / EMU, s.height / EMU
            if MIN_ICON_IN <= max(w, h) <= 1.6 and min(w, h) > 0 and 0.6 <= w / h <= 1.6 and s.shape_type is not None \
                    and any(k in str(s.shape_type) for k in ("PICTURE", "GROUP", "FREEFORM")):
                m = re.match(r"Icon - (.+)", s.name or "")
                found.append({"slide": n, "shape_id": s.shape_id, "size": round(max(w, h), 2),
                              "icon_id": m.group(1) if m else None, "color_slot": colors.icon_slot(theme, s)})
    return found


def detect_settings(prs: Presentation, library, existing: list[dict]) -> DeckSettings:
    """Fallback order: library style guide > deck's existing icons > defaults."""
    sg, sources = library.style_guide, {}
    tiers = dict(config.DEFAULT_TIERS)
    scale = (prs.slide_width / EMU) / 13.333
    tiers = {k: round(v * scale, 2) for k, v in tiers.items()}
    sources["tiers"] = "default starting points (scaled to slide width)"
    if existing:
        # TODO: assign existing sizes to tiers by placement kind; for now the most common size = card tier
        common = Counter(e["size"] for e in existing).most_common(1)[0][0]
        tiers["card"] = tiers["column"] = common
        sources["tiers"] = "deck's existing icons (card tier) + defaults"
    if sg.get("tiers"):
        tiers.update(sg["tiers"]); sources["tiers"] = "library style guide"

    styles = [library.get(e["icon_id"]).style for e in existing if e["icon_id"] and library.get(e["icon_id"])]
    style = sg.get("style") or (Counter(styles).most_common(1)[0][0] if styles else config.DEFAULT_STYLE)
    sources["style"] = "library style guide" if sg.get("style") else ("deck's existing icons" if styles else "default")

    # Colors always come from the theme. The deck color: the existing icons' theme color when they agree
    # (style guide, Color), else accent1. Each icon may still switch slot for contrast (colors.assign).
    used = Counter(e["color_slot"] for e in existing if e.get("color_slot"))
    if sg.get("color_slot") in THEME_CLASS:
        slot, why = sg["color_slot"], "library style guide"
    elif used and used.most_common(1)[0][1] >= AGREE * sum(used.values()):
        slot, n = used.most_common(1)[0]
        why = f"matches the deck's existing icons ({n} of {sum(used.values())} use it)"
    else:
        slot = "accent1"
        why = ("default; the deck's existing icons use mixed colors ("
               + ", ".join(f"{s} ×{k}" for s, k in used.most_common()) + ")") if used else "default"
    masters = Counter(s.slide_layout.slide_master.part.partname for s in prs.slides)
    main = next((m for m in prs.slide_masters if masters and m.part.partname == masters.most_common(1)[0][0]),
                prs.slide_masters[0])   # the master most slides use; each icon still uses its own slide's theme
    color = colors.Theme(main).slot_hex(slot)
    sources["color"] = (f"theme {slot} ({why}); an icon switches to its label's color or the theme's dark or "
                        f"light color where {slot} has under 3:1 contrast with what is behind it")
    return DeckSettings(style=style, color_hex=color, tiers=tiers, sources=sources, color_slot=slot)
