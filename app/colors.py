"""Icon colors, always from the deck's theme, chosen per icon for contrast with what sits behind it.

Every color is a theme slot (accent1, tx1, bg1, ...), resolved to hex through the slide master's color map
and theme, so the inserted SVG can be bound to the slot. What sits behind an icon is the slide background
with every filled shape under the icon's center composited on top in z-order, so a 15%-opacity card counts
as the pale tint it shows (groups included). Then, per icon set (a row of cards shares one color):
  1. the deck color, if it has at least 3:1 contrast (the WCAG minimum for graphics);
  2. else the label's own text color, if it reaches 2:1 (the designer already uses it there for text);
  3. else the theme's dark or light colors at 3:1; failing all, the best of them, flagged in the note."""
import colorsys
import math
import re
from collections import Counter, defaultdict

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .icon_tools import EMU, THEME_CLASS

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
MIN_CONTRAST = 3.0
LABEL_MIN_CONTRAST = 2.0
FALLBACK_SLOTS = ["tx1", "bg1", "tx2", "bg2"]
DEFAULT_MAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
PRESET = {"black": "000000", "white": "FFFFFF"}
IMAGE = "image"         # a picture or other fill whose color can't be read
_UNSET = object()       # no fill element at all (fall back to the shape style)


def _name(el) -> str:
    return etree.QName(el).localname


class Theme:
    """A master's theme colors and color map."""

    def __init__(self, master):
        self.scheme = {}
        theme = next((r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme")), None)
        if theme is not None:
            scheme = etree.fromstring(theme.blob).find(".//a:clrScheme", NS)
            for slot in scheme if scheme is not None else []:
                if len(slot):
                    c = slot[0]
                    self.scheme[_name(slot)] = (c.get("val") if _name(c) == "srgbClr" else c.get("lastClr", "000000")).upper()
        cmap = master._element.find("p:clrMap", NS)
        self.map = dict(cmap.attrib) if cmap is not None else dict(DEFAULT_MAP)

    def slot_hex(self, slot: str) -> str:
        return self.scheme.get(self.map.get(slot, slot), "000000")

    def slot_for_hex(self, hexv: str) -> str | None:
        """The theme slot with exactly this color, if any (black text is tx1 in most themes)."""
        return next((s for s in FALLBACK_SLOTS + [f"accent{n}" for n in range(1, 7)] if self.slot_hex(s) == hexv), None)

    def color(self, el) -> tuple[str, str | None, float]:
        """(hex, theme slot or None, opacity) for a DrawingML color element, lumMod/lumOff/tint/shade applied.
        The slot is None when the color was lightened or darkened, or matches no theme color."""
        kind, slot = _name(el), None
        if kind == "schemeClr":
            slot = el.get("val")
            if slot in ("dk1", "lt1", "dk2", "lt2"):  # rare direct use: name it the way slides do
                slot = {v: k for k, v in self.map.items()}.get(slot, slot)
            hexv = self.slot_hex(slot)
        elif kind == "srgbClr":
            hexv = el.get("val", "000000").upper()
        elif kind == "sysClr":
            hexv = el.get("lastClr", "000000").upper()
        elif kind == "prstClr":
            hexv = PRESET.get(el.get("val"), "000000")
        else:
            return "000000", None, 1.0
        alpha = next((int(t.get("val", "100000")) / 100000 for t in el if _name(t) == "alpha"), 1.0)
        if any(_name(t) in ("lumMod", "lumOff", "tint", "shade") for t in el):
            return _transform(hexv, el), None, alpha
        return hexv, slot or self.slot_for_hex(hexv), alpha


def _transform(hexv: str, el) -> str:
    r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for t in el:
        kind, v = _name(t), int(t.get("val", "100000")) / 100000
        if kind == "lumMod":
            l *= v
        elif kind == "lumOff":
            l += v
        elif kind in ("tint", "shade"):
            r, g, b = colorsys.hls_to_rgb(h, min(1, max(0, l)), s)
            r, g, b = ((c * v + 1 - v) if kind == "tint" else c * v for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
    r, g, b = colorsys.hls_to_rgb(h, min(1, max(0, l)), s)
    return "%02X%02X%02X" % tuple(round(c * 255) for c in (r, g, b))


def _blend(top: str, bottom: str, alpha: float) -> str:
    return "%02X%02X%02X" % tuple(round(int(top[i:i + 2], 16) * alpha + int(bottom[i:i + 2], 16) * (1 - alpha))
                                  for i in (0, 2, 4))


def contrast(hex_a: str, hex_b: str) -> float:
    def lum(hexv):
        c = [int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        c = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    a, b = sorted((lum(hex_a), lum(hex_b)), reverse=True)
    return (a + 0.05) / (b + 0.05)


def _fill(theme, props, group_fill=None):
    """(hex, opacity), None (no fill), IMAGE, or _UNSET for an spPr / bgPr / grpSpPr element."""
    for child in props if props is not None else []:
        kind = _name(child)
        if kind == "noFill":
            return None
        if kind == "solidFill":
            return theme.color(child[0])[::2] if len(child) else None
        if kind == "gradFill":
            stops = [theme.color(gs[0]) for gs in child.findall("a:gsLst/a:gs", NS) if len(gs)]
            return (_mix([s[0] for s in stops]), sum(s[2] for s in stops) / len(stops)) if stops else None
        if kind == "pattFill":
            fg = child.find("a:fgClr", NS)
            return theme.color(fg[0])[::2] if fg is not None and len(fg) else None
        if kind == "blipFill":
            return IMAGE
        if kind == "grpFill":
            return group_fill
    return _UNSET


def _mix(hexes):
    return "%02X%02X%02X" % tuple(round(sum(int(h[i:i + 2], 16) for h in hexes) / len(hexes)) for i in (0, 2, 4))


def _shape_fill(theme, el, group_fill):
    kind = _name(el)
    if kind == "pic":
        return IMAGE
    if kind != "sp":                       # lines, tables, charts: not treated as a background
        return None
    f = _fill(theme, el.find("p:spPr", NS), group_fill)
    if f is not _UNSET:
        return f
    ref = el.find("p:style/a:fillRef", NS)  # no fill of its own: the theme style decides
    if ref is not None and ref.get("idx", "0") != "0" and len(ref):
        return theme.color(ref[0])[::2]
    return None


def shape_fill(theme, sh):
    """A shape's own fill: (hex, opacity), None (no fill) or IMAGE."""
    return _shape_fill(theme, sh._element, None)


def is_drawn(sh, fill) -> bool:
    """Has visible edges: a picture, line, table or chart, or a shape with a fill or outline."""
    el = sh._element
    if _name(el) in ("pic", "graphicFrame", "cxnSp"):
        return True
    if fill is not None:
        return True
    ln = el.find("p:spPr/a:ln", NS)
    if ln is not None and ln.find("a:noFill", NS) is not None:
        return False
    if ln is not None and any(_name(c) in ("solidFill", "gradFill", "pattFill") for c in ln):
        return True
    ref = el.find("p:style/a:lnRef", NS)
    return ref is not None and ref.get("idx", "0") != "0"


RECTISH = {"rect", "roundRect", "snip1Rect", "snip2SameRect", "snip2DiagRect", "snipRoundRect", "round1Rect",
           "round2SameRect", "round2DiagRect", "flowChartProcess", "flowChartAlternateProcess", "line"}


def geometry(sh) -> str:
    """The shape's outline: a preset name (rect, hexagon, ellipse...), "custom", "picture", or "none"."""
    el = sh._element
    if _name(el) in ("pic", "graphicFrame"):
        return "picture" if _name(el) == "pic" else "table"
    if _name(el) == "cxnSp":
        return "line"
    prst = el.find("p:spPr/a:prstGeom", NS)
    return prst.get("prst", "rect") if prst is not None else ("custom" if el.find("p:spPr/a:custGeom", NS) is not None
                                                             else "none")


def _polygon(geom: str, rotation: float, box):
    """Corners of a preset hexagon, octagon or diamond within its on-screen box (turned a quarter at a time);
    None for anything else, whose box stands in for its outline."""
    quarters = (rotation or 0) / 90
    if abs(quarters - round(quarters)) > 0.02:
        return None
    turned = round(quarters) % 2 == 1
    l, t, r, b = box
    cx, cy, s = (l + r) / 2, (t + b) / 2, min(r - l, b - t)
    if geom == "hexagon":
        a = 0.25 * s
        return [(cx, t), (r, t + a), (r, b - a), (cx, b), (l, b - a), (l, t + a)] if turned else \
               [(l, cy), (l + a, t), (r - a, t), (r, cy), (r - a, b), (l + a, b)]
    if geom == "octagon":
        a = 0.29289 * s
        return [(l + a, t), (r - a, t), (r, t + a), (r, b - a), (r - a, b), (l + a, b), (l, b - a), (l, t + a)]
    if geom == "diamond":
        return [(cx, t), (r, cy), (cx, b), (l, cy)]
    return None


def inside_outline(geom: str, rotation: float, box, x: float, y: float) -> bool:
    """Is the point inside the shape's actual outline? Ellipses and preset polygons are tested exactly;
    rectangles, freeforms and the rest use their box."""
    l, t, r, b = box
    if geom == "ellipse":
        rx, ry = (r - l) / 2 or 1e-9, (b - t) / 2 or 1e-9
        return ((x - (l + r) / 2) / rx) ** 2 + ((y - (t + b) / 2) / ry) ** 2 <= 1.0
    poly = _polygon(geom, rotation, box)
    if poly is None:
        return l <= x <= r and t <= y <= b
    side = 0
    for (x1, y1), (x2, y2) in zip(poly, poly[1:] + poly[:1]):
        cross = (x2 - x1) * (y - y1) - (y2 - y1) * (x - x1)
        if abs(cross) > 1e-12:
            if side and (cross > 0) != (side > 0):
                return False
            side = cross
    return True


def rotated_box(x, y, w, h, degrees):
    """(left, top, width, height) of the area a shape covers on screen once rotated about its center."""
    if not degrees % 180:
        return x, y, w, h
    rad = math.radians(degrees)
    c, s = abs(math.cos(rad)), abs(math.sin(rad))
    nw, nh = w * c + h * s, w * s + h * c
    return x + (w - nw) / 2, y + (h - nh) / 2, nw, nh


def _child_tf(group, x, y, w, h, tf):
    """How a group at (x, y, w, h) on the slide maps its children's coordinates: (dx, dy, scale x, scale y)."""
    xfrm = group._element.grpSpPr.find("a:xfrm", NS)
    off = xfrm.find("a:chOff", NS) if xfrm is not None else None
    ext = xfrm.find("a:chExt", NS) if xfrm is not None else None
    if off is None or ext is None:   # no child frame: children already use the group's coordinates
        return tf
    csx, csy = w / (int(ext.get("cx")) or 1), h / (int(ext.get("cy")) or 1)
    return x - int(off.get("x")) * csx, y - int(off.get("y")) * csy, csx, csy


class Placed:
    """A shape inside a group, seen where it sits on the slide: left, top, width and height in slide EMU (group
    offsets and scaling applied, before its own rotation); `scale` is the group's horizontal scaling.
    Everything else is the shape's own."""

    def __init__(self, shape, left, top, width, height, scale=1.0):
        self._shape, self.left, self.top, self.width, self.height = shape, left, top, width, height
        self.scale = scale

    def __getattr__(self, name):
        return getattr(self._shape, name)


def leaves(shapes, tf=(0.0, 0.0, 1.0, 1.0), groups=False):
    """Every shape that isn't a group, in z-order, as placed on the slide (a Placed when inside a group);
    with groups=True, each group too (as Placed), before its children."""
    ox, oy, sx, sy = tf
    for sh in shapes:
        if sh.left is None or sh.width is None:
            continue
        if tf == (0.0, 0.0, 1.0, 1.0) and sh.shape_type != MSO_SHAPE_TYPE.GROUP:
            yield sh
            continue
        x, y, w, h = ox + sh.left * sx, oy + sh.top * sy, sh.width * sx, sh.height * sy
        placed = Placed(sh, round(x), round(y), round(w), round(h), sx)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            if groups:
                yield placed
            yield from leaves(sh.shapes, _child_tf(sh, x, y, w, h, tf), groups)
        else:
            yield placed


def find_shape(shapes, shape_id):
    """The shape with this id, inside groups too (as placed on the slide), or None."""
    return next((s for s in leaves(shapes) if s.shape_id == shape_id), None)


def boxes(theme, shapes, tf=(0.0, 0.0, 1.0, 1.0), group_fill=None):
    """(left, top, right, bottom) in EMU, fill, and shape for every shape in z-order; group children on
    slide coordinates."""
    ox, oy, sx, sy = tf
    for sh in shapes:
        if sh.left is None or sh.width is None:
            continue
        x, y, w, h = ox + sh.left * sx, oy + sh.top * sy, sh.width * sx, sh.height * sy
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            own = _fill(theme, sh._element.grpSpPr, group_fill)
            yield from boxes(theme, sh.shapes, _child_tf(sh, x, y, w, h, tf), group_fill if own is _UNSET else own)
        else:
            x, y, w, h = rotated_box(x, y, w, h, getattr(sh, "rotation", 0) or 0)
            yield (x, y, x + w, y + h), _shape_fill(theme, sh._element, group_fill), sh


def _slide_background(theme, slide):
    """Hex or IMAGE: the slide's background, else its layout's, else its master's, else bg1."""
    for part in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = part._element.find("p:cSld/p:bg", NS)
        if bg is None:
            continue
        f = _fill(theme, bg.find("p:bgPr", NS))
        if f == IMAGE:
            return IMAGE
        if f is not _UNSET and f is not None:
            return _blend(f[0], "FFFFFF", f[1])
        ref = bg.find("p:bgRef", NS)
        if ref is not None and len(ref):
            return theme.color(ref[0])[0]
    return theme.slot_hex("bg1")


def layers(slide) -> list[tuple[list, bool]]:
    """What draws on a slide, back to front, as (shapes, on the slide itself?): the master's and layout's
    decorations (lines, bands, logos; not their placeholders) unless the slide hides them, then its own shapes."""
    out = []
    layout = slide.slide_layout
    if slide._element.get("showMasterSp", "1") != "0":
        if layout._element.get("showMasterSp", "1") != "0":
            out.append(([s for s in layout.slide_master.shapes if not s.is_placeholder], False))
        out.append(([s for s in layout.shapes if not s.is_placeholder], False))
    out.append((list(slide.shapes), True))
    return out


def behind(theme, slide, x_in, y_in):
    """(hex or IMAGE, what it is) at a point: the background with every filled shape under the point
    composited on top in z-order."""
    x, y = x_in * EMU, y_in * EMU
    color, where = _slide_background(theme, slide), "the slide background"
    drawn = [entry for shapes, _ in layers(slide) for entry in boxes(theme, shapes)]
    for (l, t, r, b), fill, _ in drawn:
        if fill is None or not (l <= x <= r and t <= y <= b):
            continue
        if fill == IMAGE:
            color, where = IMAGE, "a picture"
            continue
        hexv, alpha = fill
        if alpha >= 0.999 or (color == IMAGE and alpha >= 0.5):  # mostly opaque over a picture: the fill dominates
            color = hexv
        elif color != IMAGE:
            color = _blend(hexv, color, alpha)
        where = "the shape behind it"
    return color, where


NEAR = 40               # RGB distance within which an existing icon's color counts as a theme color
CLASS_SLOT = {name: slot for slot, name in THEME_CLASS.items()}


def near_slot(theme, hexv: str) -> str | None:
    """The theme slot closest to a color, if it is close enough to be meant as that theme color."""
    rgb = [int(hexv[i:i + 2], 16) for i in (0, 2, 4)]
    def dist(slot):
        h = theme.slot_hex(slot)
        return sum((a - int(h[i:i + 2], 16)) ** 2 for a, i in zip(rgb, (0, 2, 4))) ** 0.5
    best = min(THEME_CLASS, key=dist)
    return best if dist(best) <= NEAR else None


def icon_slot(theme, shape) -> str | None:
    """The theme slot an existing icon is drawn in: its SVG's theme class, else its main color (SVG paint,
    PNG pixels, or the fills and outlines of a native icon) when that is a theme color. None otherwise."""
    from .icon_tools import NS_SVG, NS_R, svg_paint
    el = shape._element
    counts = Counter()
    if _name(el) == "pic":
        svg_blip = el.find(".//{%s}svgBlip" % NS_SVG)
        try:
            if svg_blip is not None:
                svg = shape.part.related_part(svg_blip.get("{%s}embed" % NS_R)).blob
                named = Counter(m.decode() for m in re.findall(rb"MsftOfcThm_(\w+?)_(?:Fill|Stroke)_v2", svg))
                if named:
                    return CLASS_SLOT.get(named.most_common(1)[0][0])
                counts = svg_paint(svg)
            else:
                counts = _pixel_colors(shape.image.blob)
        except (KeyError, ValueError, OSError, etree.XMLSyntaxError):
            return None
    else:  # a group or freeform drawn with shapes: the colors of its fills and outlines, not its text
        for fill in el.iter("{%s}solidFill" % NS["a"]):
            if _name(fill.getparent()) in ("spPr", "ln") and len(fill):
                hexv, slot, _ = theme.color(fill[0])
                counts[theme.slot_hex(slot) if slot else hexv] += 1
    if not counts:
        return None
    return near_slot(theme, counts.most_common(1)[0][0])


def _pixel_colors(blob) -> Counter:
    """The main colors of a raster icon: its opaque pixels, or on an opaque image the pixels that differ
    from the border (the background)."""
    from io import BytesIO
    from PIL import Image
    im = Image.open(BytesIO(blob)).convert("RGBA")
    im.thumbnail((64, 64))
    raw = im.tobytes()
    px = [tuple(raw[i:i + 4]) for i in range(0, len(raw), 4)]
    if min(p[3] for p in px) < 250:
        px = [p for p in px if p[3] >= 250]
    else:
        w, h = im.size
        edge = Counter(im.getpixel((x, y))[:3] for x in range(w) for y in (0, h - 1)).most_common(1)[0][0]
        px = [p for p in px if sum(abs(a - b) for a, b in zip(p[:3], edge)) > 60]
    return Counter("%02X%02X%02X" % p[:3] for p in px)


def text_slot(theme, slide, item) -> str | None:
    """The theme slot of the item's text color, when the slide sets one: the run's color, else the box's list
    style, else the shape style's font color. None when it's inherited from the master (not read) or isn't a
    theme color, so an assumed color never counts as the designer's choice."""
    shape = find_shape(slide.shapes, item.shape_id)
    if shape is None or not shape.has_text_frame:
        return None
    paras = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    para = paras[item.paragraph] if item.paragraph < len(paras) else (paras[0] if paras else None)
    runs = [r for r in (para.runs if para else []) if r.text.strip()]
    sources = [r._r.find("a:rPr/a:solidFill", NS) for r in runs[:1]] + [
        shape._element.find("p:txBody/a:lstStyle/a:lvl1pPr/a:defRPr/a:solidFill", NS),
        shape._element.find("p:style/a:fontRef", NS)]
    for el in sources:
        if el is not None and len(el):
            return theme.color(el[0])[1]
    return None


def assign(prs, picks, default_slot: str = "accent1") -> None:
    """Set color_slot, color_hex and color_note on each pick. Icons in one set share one color, unless the
    designer colored the set's labels differently per background: then each icon matches its own label."""
    themes = {}
    sets = defaultdict(list)
    for p in picks:
        sets[p.item.set_id or f"solo-{id(p)}"].append(p)
    for members in sets.values():
        seen = []  # (pick, background hex or IMAGE, what it is, text slot)
        for p in members:
            slide = prs.slides[p.item.slide - 1]
            master = slide.slide_layout.slide_master
            theme = themes.setdefault(id(master), Theme(master))
            bg, where = behind(theme, slide, p.left + p.size / 2, p.top + p.size / 2)
            seen.append((p, bg, where, text_slot(theme, slide, p.item)))
        known = [(bg, where) for _, bg, where, _ in seen if bg != IMAGE]
        labels = [s for *_, s in seen if s in THEME_CLASS and s != default_slot]
        candidates = [(default_slot, MIN_CONTRAST, "deck color")]
        candidates += [(s, LABEL_MIN_CONTRAST, "the label's color") for s in dict.fromkeys(labels)]
        candidates += [(s, MIN_CONTRAST, "theme dark/light") for s in FALLBACK_SLOTS if s != default_slot]

        def worst(slot):
            return min((contrast(theme.slot_hex(slot), bg), bg, where) for bg, where in known)

        # The designer gave the labels different colors for their backgrounds (white heading on a dark panel,
        # dark heading on a light one): when the deck color doesn't work for all, each icon matches its label.
        own = [s for *_, s in seen]
        if known and worst(default_slot)[0] < MIN_CONTRAST and len(set(own)) > 1 and all(s in THEME_CLASS for s in own) \
                and all(bg == IMAGE or contrast(theme.slot_hex(s), bg) >= LABEL_MIN_CONTRAST for _, bg, _, s in seen):
            for p, bg, where, s in seen:
                ratio = contrast(theme.slot_hex(s), bg) if bg != IMAGE else None
                p.color_slot, p.color_hex = s, theme.slot_hex(s)
                p.color_note = (f"{s} (matches its label): the labels in this set use different colors for their "
                                f"backgrounds" + (f"; {ratio:.1f}:1 on {where} (#{bg})" if ratio else ""))
            continue
        if not known:
            slot, note = default_slot, f"{default_slot} (deck color): sits on a picture, so contrast wasn't checked"
        else:
            chosen = next((c for c in candidates if worst(c[0])[0] >= c[1]), None)
            slot, _, why = chosen or max(candidates, key=lambda c: worst(c[0])[0])
            ratio, bg, where = worst(slot)
            if slot == default_slot:
                note = f"{slot} (deck color): {ratio:.1f}:1 on {where} (#{bg})"
            else:
                d_ratio, d_bg, d_where = worst(default_slot)
                note = (f"{slot} ({why}): the deck color {default_slot} is only {d_ratio:.1f}:1 on {d_where} "
                        f"(#{d_bg}); {slot} is {ratio:.1f}:1")
                if not chosen:
                    note += ", still low: check this slide"
        for p, *_ in seen:
            p.color_slot, p.color_hex, p.color_note = slot, theme.slot_hex(slot), note
